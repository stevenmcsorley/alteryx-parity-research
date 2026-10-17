"""Excel (.xlsx/.xlsm) and OpenDocument (.ods) worksheets for Input Data.

Excel is read with openpyxl (cached formula values, not formulas). ODS is parsed directly from the
file's content.xml, streaming, so the large "repeated empty row" runs spreadsheet programs write at
the end of a sheet are never expanded. Published statistical tables usually have notes above the
header; ``skip_rows`` (rows above the header) and ``range`` (for example ``A5:H200``) select the table.
"""
from __future__ import annotations

import datetime as _dt
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from .errors import ConfigError, DataError
from .types import Field, FieldType, format_value

SPREADSHEET_FORMATS = ('xlsx', 'ods')
UNBOUNDED_TEXT = 1073741823
_NS = {'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
       'office': 'urn:oasis:names:tc:opendocument:xmlns:office:1.0',
       'text': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'}
_T = '{%s}' % _NS['table']
_O = '{%s}' % _NS['office']
_X = '{%s}' % _NS['text']


def _cell_ref(ref: str) -> tuple[int, int]:
    m = re.fullmatch(r'([A-Za-z]{1,3})(\d+)', ref.strip())
    if not m:
        raise ConfigError(f'Invalid cell reference {ref!r}; use a form such as A5')
    col = 0
    for ch in m.group(1).upper():
        col = col * 26 + ord(ch) - 64
    return int(m.group(2)), col


def parse_range(text: str | None):
    """'A5:H200' -> (first_row, first_col, last_row, last_col), 1-based and inclusive. 'A5' leaves the end open."""
    if not text:
        return None
    parts = text.split(':')
    if len(parts) > 2:
        raise ConfigError(f'Invalid range {text!r}; use a form such as A5:H200')
    r1, c1 = _cell_ref(parts[0])
    r2, c2 = _cell_ref(parts[1]) if len(parts) == 2 else (None, None)
    if r2 is not None and (r2 < r1 or c2 < c1):
        raise ConfigError(f'Invalid range {text!r}: the end is before the start')
    return r1, c1, r2, c2


# ------------------------------------------------------------------------------------------ reading
def _xlsx_rows(path: Path, sheet: str | None):
    try:
        import openpyxl
    except ImportError:  # pragma: no cover - bundled in every runtime
        raise DataError('Reading Excel files needs openpyxl, which is missing from this installation') from None
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError:
        raise DataError(f'Input file not found: {path}') from None
    except Exception as e:
        raise DataError(f'{path.name}: cannot read as an Excel workbook: {e}') from None
    try:
        names = wb.sheetnames
        ws = wb[_pick(names, sheet, path)]
        for row in ws.iter_rows(values_only=True):
            yield list(row)
    finally:
        wb.close()


def _pick(names, sheet, path):
    if sheet is None:
        if not names:
            raise DataError(f'{path.name} has no worksheets')
        return names[0]
    if sheet in names:
        return sheet
    raise DataError(f'{path.name} has no sheet {sheet!r}; sheets are {names}')


def _ods_value(cell):
    kind = cell.get(_O + 'value-type')
    if kind in ('float', 'percentage', 'currency'):
        v = float(cell.get(_O + 'value'))
        return int(v) if v.is_integer() and abs(v) < 2 ** 53 else v
    if kind == 'date':
        raw = cell.get(_O + 'date-value')
        return _dt.datetime.fromisoformat(raw) if 'T' in raw else _dt.date.fromisoformat(raw)
    if kind == 'time':
        m = re.fullmatch(r'PT(\d+)H(\d+)M(\d+(?:\.\d+)?)S', cell.get(_O + 'time-value', ''))
        if m:
            secs = float(m.group(3))
            return _dt.time(int(m.group(1)) % 24, int(m.group(2)), int(secs), int(round((secs % 1) * 1e6)) % 1000000)
    if kind == 'boolean':
        return cell.get(_O + 'boolean-value') == 'true'
    paragraphs = [''.join(p.itertext()) for p in cell.iter(_X + 'p')]
    text = '\n'.join(paragraphs)
    return text if text or kind == 'string' else None


def _ods_sheets(path: Path):
    try:
        z = zipfile.ZipFile(path)
    except FileNotFoundError:
        raise DataError(f'Input file not found: {path}') from None
    except zipfile.BadZipFile:
        raise DataError(f'{path.name}: not an OpenDocument spreadsheet') from None
    return z


def _ods_iter(path: Path, only: str | None = None, first_only: bool = False):
    """Yields (sheet name, row) for every non-empty row (empty rows inside the data as []), in one pass."""
    z = _ods_sheets(path)
    with z, z.open('content.xml') as fh:
        name, wanted, pending_empty, seen_first = None, False, 0, False
        for event, el in ET.iterparse(fh, events=('start', 'end')):
            if el.tag == _T + 'table':
                if event == 'start':
                    name = el.get(_T + 'name')
                    wanted = (only is None or name == only) and not (first_only and seen_first)
                    seen_first, pending_empty = True, 0
                else:
                    el.clear()
                    if wanted and (only is not None or first_only):
                        return
                continue
            if not wanted or event != 'end' or el.tag != _T + 'table-row':
                continue
            row, empty_run = [], 0
            for cell in el:
                if cell.tag not in (_T + 'table-cell', _T + 'covered-table-cell'):
                    continue
                repeat = int(cell.get(_T + 'number-columns-repeated', '1'))
                value = _ods_value(cell) if cell.tag == _T + 'table-cell' else None
                if value is None:
                    empty_run += repeat
                    continue
                row.extend([None] * empty_run)
                empty_run = 0
                row.extend([value] * min(repeat, 16384))
            repeat_rows = int(el.get(_T + 'number-rows-repeated', '1'))
            el.clear()
            if not row:
                pending_empty += repeat_rows
                continue
            for _ in range(pending_empty):
                yield name, []
            pending_empty = 0
            for _ in range(min(repeat_rows, 1_048_576)):
                yield name, list(row)


def _ods_rows(path: Path, sheet: str | None):
    if sheet is not None and sheet not in list_sheet_names(path):
        raise DataError(f'{path.name} has no sheet {sheet!r}; sheets are {list_sheet_names(path)}')
    for _, row in _ods_iter(path, sheet, first_only=sheet is None):
        yield row


def list_sheet_names(path: Path) -> list[str]:
    if path.suffix.lower() == '.ods':
        z = _ods_sheets(path)
        with z, z.open('content.xml') as fh:
            return [el.get(_T + 'name') for ev, el in ET.iterparse(fh, events=('start',)) if el.tag == _T + 'table']
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def grid(path: Path, fmt: str, sheet: str | None, cell_range: str | None, limit_rows: int | None = None) -> list[list]:
    """The selected cells as rows of Python values, with trailing empty rows and columns removed."""
    rows_iter = _xlsx_rows(path, sheet) if fmt == 'xlsx' else _ods_rows(path, sheet)
    bounds = parse_range(cell_range)
    r1, c1, r2, c2 = bounds if bounds else (1, 1, None, None)
    out = []
    for number, row in enumerate(rows_iter, start=1):
        if number < r1:
            continue
        if r2 is not None and number > r2:
            break
        row = row[c1 - 1:c2] if c2 is not None else row[c1 - 1:]
        out.append([None if (isinstance(v, str) and v == '') else v for v in row])
        if limit_rows is not None and len(out) >= limit_rows:
            break
    return _trim(out)


def _trim(rows: list[list]) -> list[list]:
    """Drops trailing empty rows and pads every row to the last used column."""
    out = [[None if (isinstance(v, str) and v == '') else v for v in r] for r in rows]
    while out and all(v is None for v in out[-1]):
        out.pop()
    width = max((max((i + 1 for i, v in enumerate(r) if v is not None), default=0) for r in out), default=0)
    return [(r + [None] * width)[:width] for r in out]


def _column_type(values):
    present = [v for v in values if v is not None]
    if not present:
        return FieldType.V_WSTRING
    kinds = {type(v) for v in present}
    if kinds <= {bool}:
        return FieldType.BOOL
    if kinds <= {int}:
        return FieldType.INT64
    if kinds <= {int, float}:
        return FieldType.DOUBLE
    if kinds <= {_dt.datetime, _dt.date}:
        if all(isinstance(v, _dt.date) and (not isinstance(v, _dt.datetime) or v.time() == _dt.time()) for v in present):
            return FieldType.DATE
        return FieldType.DATETIME
    if kinds <= {_dt.time}:
        return FieldType.TIME
    return FieldType.V_WSTRING


def _convert(value, kind):
    if value is None:
        return None
    if kind is FieldType.V_WSTRING:
        return value if isinstance(value, str) else format_value(value)
    if kind is FieldType.DATE and isinstance(value, _dt.datetime):
        return value.date()
    if kind is FieldType.DATETIME and not isinstance(value, _dt.datetime):
        return _dt.datetime.combine(value, _dt.time())
    if kind is FieldType.DOUBLE:
        return float(value)
    return value


def read_table(path: Path, fmt: str, cfg: dict, warn=None):
    """-> (names, columns, natural fields) for the selected sheet and cell range."""
    cells = grid(path, fmt, cfg.get('sheet'), cfg.get('range'))
    cells = cells[cfg.get('skip_rows', 0):]
    if cfg.get('has_header', True):
        if not cells:
            raise DataError(f'{path.name}: the selected cells are empty; a header row was expected')
        header, body = cells[0], cells[1:]
        names, seen, renamed = [], {}, []
        for i, v in enumerate(header):
            name = (v if isinstance(v, str) else format_value(v) if v is not None else '').strip().replace('\n', ' ')
            if not name:
                name = f'Field_{i + 1}'
                renamed.append(name)
            if name in seen:
                seen[name] += 1
                renamed.append(f'{name}_{seen[name]}')
                name = f'{name}_{seen[name]}'
            else:
                seen[name] = 1
            names.append(name)
        if renamed and warn:
            warn(f'{len(renamed)} header cell(s) were empty or repeated and were named {renamed[:5]}', count=len(renamed))
    else:
        width = len(cells[0]) if cells else 0
        names, body = [f'Field_{i + 1}' for i in range(width)], cells
    body = [r for r in body if any(v is not None for v in r)]
    if cfg.get('max_records') is not None:
        body = body[:cfg['max_records']]
    cols = [[r[i] for r in body] for i in range(len(names))]
    fields, out = [], []
    for name, col in zip(names, cols):
        kind = _column_type(col)
        fields.append(Field(name, kind, UNBOUNDED_TEXT if kind is FieldType.V_WSTRING else None))
        out.append([_convert(v, kind) for v in col])
    return names, out, fields


def describe(path: Path, preview_rows: int = 12) -> list[dict]:
    """Every sheet with its used size and first rows as text, to choose sheet, skip_rows and range."""
    fmt = 'ods' if path.suffix.lower() == '.ods' else 'xlsx'
    by_sheet = {name: [] for name in list_sheet_names(path)}
    if fmt == 'ods':
        for name, row in _ods_iter(path):
            by_sheet[name].append(row)
    sheets = []
    for name in by_sheet:
        cells = _trim(by_sheet[name]) if fmt == 'ods' else grid(path, fmt, name, None)
        sheets.append({'sheet': name, 'rows': len(cells), 'columns': len(cells[0]) if cells else 0,
                       'first_rows': [[None if v is None else format_value(v) if not isinstance(v, str) else v for v in r]
                                      for r in cells[:preview_rows]]})
    return sheets
